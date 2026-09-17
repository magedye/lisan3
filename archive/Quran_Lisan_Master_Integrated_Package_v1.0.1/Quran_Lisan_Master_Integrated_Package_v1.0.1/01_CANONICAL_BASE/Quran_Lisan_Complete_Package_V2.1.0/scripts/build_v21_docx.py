from pathlib import Path
import re,csv
from docx import Document
from docx.shared import Inches,Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_SECTION

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'Quran_Lisan_Complete_Package_V2.1.0.docx'

def set_bidi(p,bidi=True):
    pPr=p._p.get_or_add_pPr(); el=pPr.find(qn('w:bidi'))
    if el is None:
        el=OxmlElement('w:bidi'); pPr.append(el)
    el.set(qn('w:val'),'1' if bidi else '0')

def style_run(r,font='Noto Naskh Arabic',size=10,bold=False,italic=False):
    r.font.name=font; r.font.size=Pt(size); r.bold=bold; r.italic=italic
    rPr=r._r.get_or_add_rPr(); rf=rPr.rFonts
    if rf is None:
        rf=OxmlElement('w:rFonts'); rPr.insert(0,rf)
    for k in ('w:ascii','w:hAnsi','w:cs','w:eastAsia'): rf.set(qn(k),font)

def para(doc,text='',bold=False,size=10,align=WD_ALIGN_PARAGRAPH.RIGHT,bidi=True,space_after=3):
    p=doc.add_paragraph(); p.alignment=align; set_bidi(p,bidi); p.paragraph_format.space_after=Pt(space_after); p.paragraph_format.line_spacing=1.05
    r=p.add_run(text); style_run(r,size=size,bold=bold); return p

def add_heading(doc,text,level):
    p=doc.add_paragraph(); p.style=f'Heading {min(level,4)}'; p.alignment=WD_ALIGN_PARAGRAPH.RIGHT; set_bidi(p,True); p.paragraph_format.keep_with_next=True
    r=p.add_run(text); style_run(r,size={1:17,2:14,3:12,4:11}.get(level,10),bold=True); return p

def add_code(doc,text):
    for line in text.splitlines() or ['']:
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.LEFT; set_bidi(p,False); p.paragraph_format.space_after=Pt(0); p.paragraph_format.left_indent=Inches(0.25)
        r=p.add_run(line); style_run(r,font='DejaVu Sans Mono',size=7.5)

def clean_inline(s):
    s=s.replace('**','').replace('__','').replace('<br>',' / ')
    s=re.sub(r'\[([^\]]+)\]\(([^\)]+)\)',r'\1 — \2',s)
    s=re.sub(r'`([^`]+)`',lambda m:'\u202A'+m.group(1)+'\u202C',s)
    return s

def add_table(doc,header,rows):
    if not header: return
    tbl=doc.add_table(rows=1,cols=len(header)); tbl.style='Table Grid'; tbl.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,h in enumerate(header):
        c=tbl.rows[0].cells[i]; c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.RIGHT; set_bidi(p,True); p.clear(); r=p.add_run(clean_inline(h)); style_run(r,size=8.2,bold=True)
    trPr=tbl.rows[0]._tr.get_or_add_trPr(); hdr=OxmlElement('w:tblHeader'); hdr.set(qn('w:val'),'true'); trPr.append(hdr)
    for row in rows:
        rr=tbl.add_row(); trPr=rr._tr.get_or_add_trPr(); cs=OxmlElement('w:cantSplit'); trPr.append(cs)
        row=list(row)
        for i,val in enumerate((row+['']*len(header))[:len(header)]):
            p=rr.cells[i].paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.RIGHT; set_bidi(p,True); p.clear(); r=p.add_run(clean_inline(str(val))); style_run(r,size=7.2)
    return tbl

def md_to_doc(doc,path,display_title=None):
    para(doc,f'المصدر داخل الحزمة: {path.relative_to(ROOT).as_posix()}',size=7.5,align=WD_ALIGN_PARAGRAPH.LEFT,bidi=False)
    lines=path.read_text(encoding='utf-8').splitlines(); i=0; code=False; buf=[]
    while i<len(lines):
        line=lines[i]
        if line.strip().startswith('```'):
            if not code: code=True; buf=[]
            else: add_code(doc,'\n'.join(buf)); code=False; buf=[]
            i+=1; continue
        if code: buf.append(line); i+=1; continue
        if line.startswith('|') and i+1<len(lines) and re.match(r'^\s*\|?\s*:?-+',lines[i+1]):
            hdr=[x.strip() for x in line.strip('|').split('|')]; i+=2; rows=[]
            while i<len(lines) and lines[i].startswith('|'):
                rows.append([x.strip() for x in lines[i].strip('|').split('|')]); i+=1
            add_table(doc,hdr,rows); continue
        m=re.match(r'^(#{1,6})\s+(.*)',line)
        if m:
            level=min(len(m.group(1))+1,4); add_heading(doc,clean_inline(m.group(2).strip()),level); i+=1; continue
        if line.startswith('> '):
            p=para(doc,clean_inline(line[2:].strip()),bold=True,size=10); p.paragraph_format.right_indent=Inches(0.25); i+=1; continue
        if line.startswith('- '):
            p=para(doc,'• '+clean_inline(line[2:].strip()),size=9.5); p.paragraph_format.right_indent=Inches(0.18); i+=1; continue
        if re.match(r'^\d+\.\s+',line):
            p=para(doc,clean_inline(line.strip()),size=9.5); p.paragraph_format.right_indent=Inches(0.18); i+=1; continue
        if line.strip() in ('---','***'):
            i+=1; continue
        if not line.strip():
            doc.add_paragraph(); i+=1; continue
        para(doc,clean_inline(line),size=9.6); i+=1
    doc.add_page_break()

D=Document(); sec=D.sections[0]
sec.top_margin=Inches(0.55); sec.bottom_margin=Inches(0.55); sec.left_margin=Inches(0.55); sec.right_margin=Inches(0.55)
for st in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Heading 4']:
    s=D.styles[st]; s.font.name='Noto Naskh Arabic'; s.font.size=Pt(10)
# cover
p=D.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; set_bidi(p,True); r=p.add_run('لسان القرآن'); style_run(r,size=25,bold=True)
p=D.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; set_bidi(p,True); r=p.add_run('الحزمة الشاملة المرجعية والتنفيذية'); style_run(r,size=18,bold=True)
p=D.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; set_bidi(p,True); r=p.add_run('Quran Lisan Complete Package — V2.1.0'); style_run(r,font='Noto Sans Arabic',size=16,bold=True)
for s in ['دستور حوكمة الذكاء الاصطناعي + منهج البحث + عقود التنفيذ + عقود المخرجات','الحالة: V2.1_CANONICAL_COMPREHENSIVE_BASELINE','التاريخ: 2026-09-05']:
    p=D.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; set_bidi(p,True); r=p.add_run(s); style_run(r,size=11)
D.add_page_break()
add_heading(D,'خريطة الوثيقة الموحدة',1)
para(D,'هذه النسخة مخصصة للقراءة والمراجعة. العقود الآلية الأصلية تبقى ملفات Pydantic/JSON Schema/OpenAPI/CSV داخل ZIP، وهي المرجع التنفيذي في المجالات القابلة للآلة.',size=10)
sections=[
('A','سلطة الحزمة والبدء'),('B','المرجع المنهجي والمصطلحي'),('C','مبرر اعتماد الذكاء الاصطناعي ودستور حوكمته'),('D','النموذج الدلالي وعقد المفهوم الجذري'),('E','المصادر والمعمارية والمتطلبات'),('F','المخرجات والثقة والتقييم'),('G','عقود البحث V2'),('H','إغلاق التنفيذ والعقود'),('I','Prototype وCAP-GO'),('J','ملخص السجلات الحاكمة'),('K','سجل اعتماد حوكمة AI')]
add_table(D,['القسم','المحتوى'],sections); D.add_page_break()

order=[]
for rel in ['00_PACKAGE_AUTHORITY.md','01_PACKAGE_MAP_AND_START_HERE.md','02_V2_CHANGELOG_AND_INTEGRATION.md','03_V2_ACCEPTANCE_AND_LIMITS.md','04_CANONICAL_COMPONENTS.md']:
    order.append(rel)
# Canonical reference order emphasizing new governance
for rel in [
'reference/00_README_AND_PACKAGE_AUTHORITY.md','reference/01_PRODUCT_CHARTER_AND_SCOPE.md','reference/02_TERMINOLOGY_AND_LANGUAGE_GOVERNANCE.md','reference/03_METHODOLOGY_AND_EVIDENCE_CONSTITUTION.md',
'reference/07A_AI_ADOPTION_RATIONALE_AND_EXPECTED_ROLE.md','sources/AI_CAPABILITY_EVIDENCE_NOTES.md','reference/07B_AI_GOVERNANCE_CONSTITUTION.md','reference/07C_AI_GOVERNANCE_CONTROL_MAPPING.md','reference/07_AI_CONTROL_GOVERNANCE.md',
'reference/04_SEMANTIC_ANALYSIS_MODEL.md','reference/05_ROOT_CONCEPT_CONTRACT.md','reference/06_SOURCE_ISOLATION_AND_EXTERNAL_COMPARISON.md','reference/08_TARGET_LOGICAL_ARCHITECTURE.md','reference/09_AGENTS_SKILLS_TOOLS.md','reference/10_FUNCTIONAL_REQUIREMENTS.md','reference/11_NON_FUNCTIONAL_REQUIREMENTS.md','reference/12_DATA_MEMORY_AND_TRACEABILITY_MODEL.md','reference/13_CONFIDENCE_COMPLETENESS_VERIFICATION.md','reference/14_OUTPUT_AND_RESEARCH_UX_CONTRACT.md','reference/15_TESTING_EVALUATION_ACCEPTANCE.md','reference/16_IMPLEMENTATION_STANDARDS_AND_ROADMAP.md','reference/17_DECISIONS_AND_OPEN_QUESTIONS.md','reference/18_REUSE_KIT_RECONCILIATION_OUTCOME.md','reference/19_SOURCE_TRACEABILITY.md','reference/21_IMPLEMENTATION_PROFILE_V1.md','reference/22_GEM_OPAL_CAPABILITY_GATE.md','reference/APPENDIX_A_EXPERIMENTAL_DATASETS.md']:
    if (ROOT/rel).exists(): order.append(rel)
for p in sorted((ROOT/'research').glob('*.md')): order.append(p.relative_to(ROOT).as_posix())
for p in sorted((ROOT/'implementation').glob('*.md')): order.append(p.relative_to(ROOT).as_posix())
for rel in ['prototype/PROTOTYPE_ROOT_DECISION.md','cap-go-01/PROTOCOL.md','cap-go-01/REQUALIFICATION_POLICY.md','history/AI_GOVERNANCE_ADOPTION_RECONCILIATION_V2.1.md']:
    if (ROOT/rel).exists(): order.append(rel)

for rel in order: md_to_doc(D,ROOT/rel)

# Registries summary
add_heading(D,'ملخص السجلات الحاكمة',1)
for rel,title,cols,limit in [
('registries/REQUIREMENTS_REGISTRY.csv','المتطلبات',['id','title','priority','prototype_priority','status','version'],95),
('registries/CONTROL_REGISTRY.csv','الضوابط',['id','name','type','scope','status','version'],52),
('registries/DECISION_REGISTER.csv','القرارات',['id','decision','status','version'],47),
('registries/TERMINOLOGY_REGISTRY.csv','المصطلحات',['id','term_ar','definition','status','version'],68),
('acceptance/TEST_CATALOG.csv','اختبارات القبول',['test_id','class','title','pass_criterion','version'],47),
('sources/AI_CAPABILITY_EVIDENCE_REGISTRY.csv','أدلة قدرات AI الخارجية',['evidence_id','example','verified_claim','project_design_relevance','explicit_limit','evidence_status'],20),
]:
    para(D,title,bold=True,size=12)
    with (ROOT/rel).open(encoding='utf-8-sig',newline='') as f: rows=list(csv.DictReader(f))
    add_table(D,cols,[[r.get(c,'') for c in cols] for r in rows[:limit]])
    D.add_paragraph()
D.add_page_break()

add_heading(D,'ملخص العقود الآلية',1)
para(D,'العقود المصدرية في contracts/python/quran_lisan_contracts/، والمخططات المولدة في contracts/schemas/. في V2.1 أصبح SemanticDecision حكمًا بحثيًا صريحًا، وأضيف CanonicalAuthorization للاعتماد الرسمي عند الحاجة.',size=10)
with (ROOT/'contracts/tools/TOOL_REGISTRY.csv').open(encoding='utf-8-sig',newline='') as f: rows=list(csv.DictReader(f))
add_table(D,['tool_id','name','request_model','response_model','allowed_phases','side_effect'],[[r.get(k,'') for k in ['tool_id','name','request_model','response_model','allowed_phases','side_effect']] for r in rows])
para(D,f"عدد JSON Schemas المولدة داخل الحزمة: {len(list((ROOT/'contracts/schemas').glob('*.schema.json')))}.",size=10)
# footer
for sec in D.sections:
    p=sec.footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; set_bidi(p,True); r=p.add_run('Quran Lisan Complete Package — V2.1.0'); style_run(r,size=7.5)
D.save(OUT)
print(OUT)
