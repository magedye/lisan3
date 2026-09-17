from pathlib import Path
import re,csv,json,sys
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_SECTION
ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'Quran_Lisan_Implementation_Closure_Pack_v1.0.1.docx'

def set_bidi(p,on=True):
    pPr=p._p.get_or_add_pPr(); tag=pPr.find(qn('w:bidi'))
    if on and tag is None:
        tag=OxmlElement('w:bidi'); pPr.append(tag)
    if not on and tag is not None: pPr.remove(tag)

def style_run(run,font='Noto Naskh Arabic',size=11,bold=False):
    run.font.name=font; run.font.size=Pt(size); run.bold=bold
    rPr=run._r.get_or_add_rPr(); rFonts=rPr.rFonts
    if rFonts is None:
        rFonts=OxmlElement('w:rFonts'); rPr.insert(0,rFonts)
    rFonts.set(qn('w:ascii'),font); rFonts.set(qn('w:hAnsi'),font); rFonts.set(qn('w:cs'),font)

def add_ar_para(doc,text='',style=None,bold=False,size=11):
    p=doc.add_paragraph(style=style); p.alignment=WD_ALIGN_PARAGRAPH.RIGHT; set_bidi(p,True)
    r=p.add_run(text); style_run(r,size=size,bold=bold); return p

def add_code(doc,text):
    for line in text.splitlines() or ['']:
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.LEFT; set_bidi(p,False)
        p.paragraph_format.left_indent=Inches(0.3); p.paragraph_format.space_after=Pt(0)
        r=p.add_run(line); style_run(r,font='DejaVu Sans Mono',size=8)

def add_table(doc,header,rows):
    tbl=doc.add_table(rows=1,cols=len(header)); tbl.alignment=WD_TABLE_ALIGNMENT.CENTER; tbl.style='Table Grid'
    for i,h in enumerate(header):
        c=tbl.rows[0].cells[i]; c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.RIGHT; set_bidi(p,True); p.clear(); r=p.add_run(h); style_run(r,size=9,bold=True)
    # repeat header and keep rows together where possible
    trPr=tbl.rows[0]._tr.get_or_add_trPr(); hdr=OxmlElement('w:tblHeader'); hdr.set(qn('w:val'),'true'); trPr.append(hdr)
    for row in rows:
        rr=tbl.add_row(); trPr=rr._tr.get_or_add_trPr(); cs=OxmlElement('w:cantSplit'); trPr.append(cs); cells=rr.cells
        for i,val in enumerate(row):
            p=cells[i].paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.RIGHT; set_bidi(p,True); p.clear(); r=p.add_run(str(val)); style_run(r,size=8)
    return tbl

def md_to_doc(doc,path):
    lines=path.read_text(encoding='utf-8').splitlines(); i=0; code=False; buf=[]
    while i<len(lines):
        line=lines[i]
        if line.strip().startswith('```'):
            if not code: code=True; buf=[]
            else: add_code(doc,'\n'.join(buf)); code=False; buf=[]
            i+=1; continue
        if code: buf.append(line); i+=1; continue
        if line.startswith('|') and i+1<len(lines) and re.match(r'^\|?\s*:?-+',lines[i+1]):
            hdr=[x.strip() for x in line.strip('|').split('|')]; i+=2; rows=[]
            while i<len(lines) and lines[i].startswith('|'):
                rows.append([x.strip() for x in lines[i].strip('|').split('|')]); i+=1
            add_table(doc,hdr,rows); continue
        m=re.match(r'^(#{1,4})\s+(.*)',line)
        if m:
            level=len(m.group(1)); text=m.group(2).strip().replace('`','').replace('**','').replace('→',' إلى '); p=doc.add_heading(level=level); p.alignment=WD_ALIGN_PARAGRAPH.RIGHT; set_bidi(p,True); r=p.add_run(text) if not p.text else None
            # add_heading already created run if text passed? we passed no text
            if r: style_run(r,size={1:17,2:14,3:12,4:11}[level],bold=True)
            i+=1; continue
        if line.startswith('- '):
            p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.RIGHT; set_bidi(p,True); p.paragraph_format.right_indent=Inches(0.2)
            bullet=re.sub(r'`([^`]+)`', lambda m: '\u202A'+m.group(1)+'\u202C', line[2:].strip().replace('**','')); r=p.add_run('• '+bullet); style_run(r,size=10); i+=1; continue
        if re.match(r'^\d+\.\s+',line):
            p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.RIGHT; set_bidi(p,True); p.paragraph_format.right_indent=Inches(0.2)
            num=re.sub(r'`([^`]+)`', lambda m: '\u202A'+m.group(1)+'\u202C', line.strip().replace('**','')); r=p.add_run(num); style_run(r,size=10); i+=1; continue
        if not line.strip(): doc.add_paragraph(); i+=1; continue
        # inline code/backticks left as text, strip markdown bold
        text=re.sub(r'`([^`]+)`', lambda m: '\u202A'+m.group(1)+'\u202C', line.replace('**',''))
        add_ar_para(doc,text,size=10); i+=1

D=Document()
sec=D.sections[0]; sec.top_margin=Inches(0.7); sec.bottom_margin=Inches(0.7); sec.left_margin=Inches(0.65); sec.right_margin=Inches(0.65)
styles=D.styles
for st in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Heading 4']:
    s=styles[st]; s.font.name='Noto Naskh Arabic'; s.font.size=Pt(11)
# cover
p=D.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; set_bidi(p,True)
r=p.add_run('لسان القرآن'); style_run(r,size=24,bold=True)
p=D.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; set_bidi(p,True)
r=p.add_run('Quran Lisan Implementation Closure Pack'); style_run(r,font='Noto Sans Arabic',size=18,bold=True)
p=D.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; set_bidi(p,True)
r=p.add_run('الحزمة التنفيذية لإغلاق المواصفات — v1.0.1'); style_run(r,size=16,bold=True)
for t in ['الحالة: IMPLEMENTATION_CLOSURE_VERIFIED','المرجع الأب: Quran Lisan Final Release Package v1.0.0','التاريخ: 2026-09-05']:
    p=D.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; set_bidi(p,True); r=p.add_run(t); style_run(r,size=11)
D.add_page_break()
# contents
add_ar_para(D,'خريطة الوثيقة',bold=True,size=18)
sections=[
('00','سلطة الإصدار'),('01','قرارات الإغلاق'),('02','نطاق التنفيذ'),('03','سياسة المصادر وFixtures'),('04','عقود البيانات'),('05','Skill/Agent/Tool mapping'),('06','إعادة إنتاج الإصدار'),('07','ابدأ من هنا'),('08','Requirement/Test/Acceptance'),('09','Migration والتوافق'),('10','البوابات المؤجلة'),('11','سجل التغيير'),('12','معيار الإغلاق'),('13','تقرير التحقق النهائي'),('Prototype','قرار PROTOTYPE-SLM-01'),('CAP-GO','بروتوكول CAP-GO-01'),('Registries','ملخص السجلات التنفيذية'),('Contracts','ملخص العقود القابلة للتنفيذ')]
add_table(D,['القسم','المحتوى'],sections)
D.add_page_break()
files=[f'{i:02d}_{n}.md' for i,n in []]
order=['00_RELEASE_AUTHORITY.md','01_CLOSURE_DECISIONS.md','02_IMPLEMENTATION_SCOPE.md','03_SOURCE_AND_FIXTURE_POLICY.md','04_DATA_CONTRACTS_GUIDE.md','05_SKILL_AGENT_TOOL_MAPPING.md','06_RELEASE_REPRODUCIBILITY.md','07_IMPLEMENTER_START_HERE.md','08_REQUIREMENT_TEST_ACCEPTANCE_GUIDE.md','09_MIGRATION_AND_COMPATIBILITY.md','10_KNOWN_DEFERRED_GATES.md','11_CHANGELOG_v1.0.0_to_v1.0.1.md','12_CLOSURE_ACCEPTANCE.md','13_FINAL_CLOSURE_VERIFICATION.md','prototype/PROTOTYPE_ROOT_DECISION.md','cap-go-01/PROTOCOL.md','cap-go-01/REQUALIFICATION_POLICY.md']
for idx,rel in enumerate(order):
    md_to_doc(D,ROOT/rel)
    if idx<len(order)-1: D.add_paragraph()
# human summaries of canonical registries
D.add_page_break(); add_ar_para(D,'ملخص السجلات التنفيذية الحاكمة',bold=True,size=18)
for rel,title,cols,limit in [
('registries/REQUIREMENTS_REGISTRY.csv','المتطلبات',['id','title','priority','prototype_priority','status'],30),
('registries/CONTROL_REGISTRY.csv','الضوابط',['id','name','type','scope','status'],25),
('registries/DECISION_REGISTER.csv','القرارات',['id','decision','status'],20),
('acceptance/TEST_CATALOG.csv','اختبارات القبول',['test_id','class','title','pass_criterion'],35)]:
    add_ar_para(D,title,bold=True,size=13)
    with (ROOT/rel).open(encoding='utf-8-sig',newline='') as f: rows=list(csv.DictReader(f))
    data=[[r.get(c,'') for c in cols] for r in rows[:limit]]
    add_table(D,cols,data)
    if len(rows)>limit: add_ar_para(D,f'يعرض هذا الملخص {limit} من {len(rows)} سجلًا؛ الملف CSV الكامل هو المرجع التنفيذي.',size=9)
# contract summary
D.add_page_break(); add_ar_para(D,'ملخص العقود القابلة للتنفيذ',bold=True,size=18)
add_ar_para(D,'العقود المصدرية محفوظة في contracts/python/quran_lisan_contracts، والمخططات المشتقة في contracts/schemas. أُدرج هذا الملخص للقراءة فقط.',size=10)
with (ROOT/'contracts/tools/TOOL_REGISTRY.csv').open(encoding='utf-8-sig',newline='') as f: rows=list(csv.DictReader(f))
add_table(D,['tool_id','name','request_model','response_model','allowed_phases','side_effect'],[[r[k] for k in ['tool_id','name','request_model','response_model','allowed_phases','side_effect']] for r in rows])
add_ar_para(D,'Skills المعتمدة: SK-01..SK-13. Agents المعتمدة: AG-01..AG-05. JSON Schemas المولدة: 48 في بيئة إنشاء الحزمة.',size=10)
# footer
for sec in D.sections:
    p=sec.footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; set_bidi(p,True); r=p.add_run('Quran Lisan — Implementation Closure Pack v1.0.1'); style_run(r,size=8)
D.save(OUT)
print(OUT)
